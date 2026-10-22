#!/usr/bin/env python3
"""Build the MLflow integration pitch deck using the Projections template."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Colors from the template
RED = RGBColor(0xEE, 0x00, 0x00)
DARK = RGBColor(0x15, 0x15, 0x15)
GRAY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xD4, 0xD4, 0xD4)
BG_DARK = RGBColor(0x1A, 0x1A, 0x1A)
BG_LIGHT = RGBColor(0xF9, 0xFA, 0xFB)
GREEN = RGBColor(0x10, 0xB9, 0x81)

FONT = "Red Hat Display"
FONT_MEDIUM = "Red Hat Display Medium"

SLD_W = Emu(12192000)
SLD_H = Emu(6858000)

def in_to_emu(inches):
    return Emu(int(inches * 914400))

def add_textbox(slide, left, top, width, height, text, font_name=FONT,
                font_size=Pt(18), color=DARK, bold=False, alignment=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf.vertical_anchor = anchor
    except:
        pass
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = font_size
    run.font.color.rgb = color
    run.font.bold = bold
    return tf

def add_multi_text(slide, left, top, width, height, lines):
    """lines: list of (text, font_name, font_size, color, bold)"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for i, (text, fn, fs, c, b) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        run = p.add_run()
        run.text = text
        run.font.name = fn
        run.font.size = fs
        run.font.color.rgb = c
        run.font.bold = b
        p.space_after = Pt(6)
    return tf

def add_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(1, left, top, width, height)  # MSO_SHAPE.RECTANGLE
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape

def build_deck():
    prs = Presentation('/Users/jkershaw/Downloads/Projections-Enterprise-Memory-Deck.pptx')

    # Remove all existing slides
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        if rId is None:
            rId = prs.slides._sldIdLst[0].attrib.get('r:id')
        prs.part.drop_rel(rId)
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])

    # ── SLIDE 1: Title ──
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.placeholders):
        sp = shape._element
        sp.getparent().remove(sp)

    add_rect(slide, 0, 0, SLD_W, SLD_H, BG_DARK)

    add_textbox(slide, in_to_emu(1.5), in_to_emu(1.2), in_to_emu(10), in_to_emu(0.8),
                "Immutable Ledger", FONT, Pt(14), LIGHT_GRAY)

    add_textbox(slide, in_to_emu(1.5), in_to_emu(2.0), in_to_emu(10), in_to_emu(1.5),
                "Cryptographic Provenance\nfor MLflow Registries",
                FONT, Pt(44), WHITE, bold=True)

    add_textbox(slide, in_to_emu(1.5), in_to_emu(4.0), in_to_emu(8), in_to_emu(0.6),
                "A proof layer for model lifecycle attestation",
                FONT, Pt(20), LIGHT_GRAY)

    add_textbox(slide, in_to_emu(1.5), in_to_emu(5.5), in_to_emu(6), in_to_emu(0.5),
                "Jonathan Kershaw  •  August 2026", FONT, Pt(12), GRAY)

    # ── SLIDE 2: The Gap ──
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.placeholders):
        sp = shape._element
        sp.getparent().remove(sp)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(0.3), in_to_emu(5), in_to_emu(0.5),
                "The Gap", FONT_MEDIUM, Pt(14), GRAY)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(0.9), in_to_emu(11), in_to_emu(1.2),
                "MLflow has no cryptographic provenance.",
                FONT, Pt(36), RED, bold=True)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(2.2), in_to_emu(10.5), in_to_emu(1.5),
                "Models move from experiment to production with metadata but no tamper evidence. "
                "There is no content hashing, no signing, no audit chain linking a registered model "
                "back to its training run with cryptographic proof. "
                "If a model artifact is modified after registration, nothing detects it.",
                FONT, Pt(16), GRAY)

    # Three column callouts
    cols = [
        ("No content hashing", "Artifacts are stored by path,\nnot by verified digest"),
        ("No lifecycle chain", "Stage transitions have metadata\nbut no tamper-evident log"),
        ("No cross-system proof", "Promotion decisions can't be\nverified outside MLflow"),
    ]
    for j, (title, desc) in enumerate(cols):
        x = in_to_emu(0.8 + j * 3.9)
        add_rect(slide, x, in_to_emu(4.0), in_to_emu(3.5), in_to_emu(2.5), BG_LIGHT)
        add_textbox(slide, x + in_to_emu(0.3), in_to_emu(4.3), in_to_emu(2.9), in_to_emu(0.5),
                    title, FONT, Pt(16), DARK, bold=True)
        add_textbox(slide, x + in_to_emu(0.3), in_to_emu(5.0), in_to_emu(2.9), in_to_emu(1.2),
                    desc, FONT, Pt(13), GRAY)

    # ── SLIDE 3: Integration Architecture ──
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.placeholders):
        sp = shape._element
        sp.getparent().remove(sp)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(0.3), in_to_emu(5), in_to_emu(0.5),
                "Integration Architecture", FONT_MEDIUM, Pt(14), GRAY)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(0.9), in_to_emu(11), in_to_emu(0.8),
                "Three layers. Zero MLflow code changes for layer one.",
                FONT, Pt(30), DARK, bold=True)

    layers = [
        ("1", "Webhook Listener", "MLflow webhook fires on model_version.created,\nalias assignments, tag changes. Listener writes\nproof receipts to the ledger. No MLflow changes.", RED, "Ready now"),
        ("2", "Artifact Wrapper", "ArtifactRepository decorator computes SHA-256\non log_artifact and records content hash to\nthe ledger. MLflow plugin entry point.", DARK, "Phase 2"),
        ("3", "Registry Wrapper", "AbstractStore wrapper intercepts create_model_version,\ntransition_stage, set_alias. Full provenance chain\nfrom experiment to production.", DARK, "Phase 3"),
    ]
    for j, (num, title, desc, accent, phase) in enumerate(layers):
        y = in_to_emu(2.1 + j * 1.7)
        add_rect(slide, in_to_emu(0.8), y, in_to_emu(0.08), in_to_emu(1.4), accent)
        add_textbox(slide, in_to_emu(1.2), y, in_to_emu(2.5), in_to_emu(0.4),
                    title, FONT, Pt(18), DARK, bold=True)
        add_textbox(slide, in_to_emu(1.2), y + in_to_emu(0.4), in_to_emu(6), in_to_emu(1.0),
                    desc, FONT, Pt(12), GRAY)
        phase_color = GREEN if phase == "Ready now" else GRAY
        add_textbox(slide, in_to_emu(10.5), y + in_to_emu(0.1), in_to_emu(2), in_to_emu(0.4),
                    phase, FONT, Pt(13), phase_color, bold=True, alignment=PP_ALIGN.RIGHT)

    # ── SLIDE 4: What We Tested ──
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.placeholders):
        sp = shape._element
        sp.getparent().remove(sp)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(0.3), in_to_emu(5), in_to_emu(0.5),
                "What We Tested", FONT_MEDIUM, Pt(14), GRAY)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(0.9), in_to_emu(11), in_to_emu(0.8),
                "Five lifecycle events. All chains verified.",
                FONT, Pt(30), DARK, bold=True)

    # Test results
    events = [
        ("Register model", "mlflow.registered_model.created"),
        ("Create model version v1", "mlflow.model_version.created"),
        ("Create model version v2", "mlflow.model_version.created"),
        ("Promote v2 to champion", "mlflow.model_version_alias.created"),
        ("Tag v2 as approved", "mlflow.model_version_tag.set"),
    ]
    add_rect(slide, in_to_emu(0.8), in_to_emu(2.0), in_to_emu(11.5), in_to_emu(3.6),
             RGBColor(0xF3, 0xF4, 0xF6))

    add_textbox(slide, in_to_emu(1.2), in_to_emu(2.2), in_to_emu(4), in_to_emu(0.4),
                "Smoke Test Results", FONT, Pt(14), DARK, bold=True)

    for j, (event, entry_type) in enumerate(events):
        y = in_to_emu(2.7 + j * 0.55)
        add_textbox(slide, in_to_emu(1.2), y, in_to_emu(0.4), in_to_emu(0.4),
                    "✓", FONT, Pt(16), GREEN, bold=True)
        add_textbox(slide, in_to_emu(1.7), y, in_to_emu(3.5), in_to_emu(0.4),
                    event, FONT, Pt(14), DARK)
        add_textbox(slide, in_to_emu(5.5), y, in_to_emu(5), in_to_emu(0.4),
                    entry_type, FONT, Pt(12), GRAY)
        add_textbox(slide, in_to_emu(10.5), y, in_to_emu(1.5), in_to_emu(0.4),
                    "proof receipt", FONT, Pt(11), GREEN, alignment=PP_ALIGN.RIGHT)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(5.8), in_to_emu(11.5), in_to_emu(1.0),
                "Each event produces a tamper-evident proof receipt with SHA-256 hash chain. "
                "Independent chain verification confirmed all entries cryptographically intact. "
                "10 unit tests + live integration test against running ledger.",
                FONT, Pt(13), GRAY)

    # ── SLIDE 5: What This Enables ──
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.placeholders):
        sp = shape._element
        sp.getparent().remove(sp)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(0.3), in_to_emu(5), in_to_emu(0.5),
                "What This Enables", FONT_MEDIUM, Pt(14), GRAY)

    add_textbox(slide, in_to_emu(0.8), in_to_emu(0.9), in_to_emu(11), in_to_emu(0.8),
                "Attestation for assets. Verified from training to production.",
                FONT, Pt(30), RED, bold=True)

    capabilities = [
        ("Model provenance", "Every model version linked to its training run,\nartifact hash, and promotion decision with\ncryptographic proof"),
        ("Cross-system verification", "Proof receipts travel with the model.\nAny system can verify without calling\nback to MLflow"),
        ("Compliance audit trail", "Immutable, append-only evidence chain.\nDB role enforced: no UPDATE, no DELETE.\nWho promoted what, when, and why"),
        ("Registry integrity", "Detect post-registration artifact tampering.\nContent-addressed storage proves the model\nyou deploy is the model you approved"),
    ]
    for j, (title, desc) in enumerate(capabilities):
        col = j % 2
        row = j // 2
        x = in_to_emu(0.8 + col * 5.9)
        y = in_to_emu(2.2 + row * 2.3)
        add_rect(slide, x, y, in_to_emu(5.5), in_to_emu(2.0), BG_LIGHT)
        add_textbox(slide, x + in_to_emu(0.3), y + in_to_emu(0.25), in_to_emu(4.9), in_to_emu(0.4),
                    title, FONT, Pt(16), DARK, bold=True)
        add_textbox(slide, x + in_to_emu(0.3), y + in_to_emu(0.7), in_to_emu(4.9), in_to_emu(1.1),
                    desc, FONT, Pt(12), GRAY)

    # ── SLIDE 6: Next Steps ──
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.placeholders):
        sp = shape._element
        sp.getparent().remove(sp)

    add_rect(slide, 0, 0, SLD_W, SLD_H, BG_DARK)

    add_textbox(slide, in_to_emu(1.5), in_to_emu(1.5), in_to_emu(10), in_to_emu(1.0),
                "Immutable Ledger + MLflow",
                FONT, Pt(36), WHITE, bold=True)

    add_textbox(slide, in_to_emu(1.5), in_to_emu(2.8), in_to_emu(9), in_to_emu(0.6),
                "Proof infrastructure for the model registry.",
                FONT, Pt(20), LIGHT_GRAY)

    steps = [
        "Webhook listener ready today — zero MLflow code changes",
        "Artifact hashing and registry wrapper as follow-on phases",
        "Looking for alignment with registry and attestation roadmap",
    ]
    for j, step in enumerate(steps):
        y = in_to_emu(3.8 + j * 0.6)
        add_textbox(slide, in_to_emu(1.5), y, in_to_emu(9), in_to_emu(0.5),
                    f"•  {step}", FONT, Pt(16), LIGHT_GRAY)

    add_textbox(slide, in_to_emu(1.5), in_to_emu(5.8), in_to_emu(6), in_to_emu(0.5),
                "Jonathan Kershaw  •  jkershaw@redhat.com", FONT, Pt(13), GRAY)

    out = '/Users/jkershaw/Documents/are-immutable-ledger/docs/mlflow-ledger-integration-pitch.pptx'
    prs.save(out)
    print(f"Saved: {out}")

build_deck()
